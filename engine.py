"""Deterministic procurement rules. No UI, AI or network is needed for decisions."""
from __future__ import annotations
import base64
import datetime as dt
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import io
import json
from pathlib import Path
import re
import zipfile
from pydantic import BaseModel, ConfigDict, Field

MAX_EXPANDED = 80 * 1024 * 1024
MAX_TEXT = 250_000
MAX_ITEMS = 500
MAX_PAGES = 300
REVIEW_EXCLUDED = {'approved', 'review_fingerprint', 'review_note', 'checked_at', 'checked_by'}

def c():
    import core
    return core

def decimal(value):
    result = c().number(value)
    if result is None:
        raise ValueError('A finite number is required.')
    return result

def cents(value):
    return decimal(value).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)

def canonical_hash(value):
    def stable(obj):
        if isinstance(obj,bool) or obj is None: return obj
        if isinstance(obj,(int,float,Decimal)):
            return {'__number__':format(Decimal(str(obj)).normalize(),'f')}
        if isinstance(obj,dict): return {key:stable(v) for key,v in obj.items()}
        if isinstance(obj,(tuple,list)): return [stable(v) for v in obj]
        return obj
    return hashlib.sha256(json.dumps(stable(value), sort_keys=True, ensure_ascii=False, allow_nan=False).encode()).hexdigest()

def field_fingerprint(field):
    return canonical_hash({k:v for k,v in c().dump(field).items() if k not in REVIEW_EXCLUDED})

def invalidate_changed(field):
    if field.approved and field.review_fingerprint != field_fingerprint(field):
        field.approved = False
        field.review_fingerprint = ''
        field.checked_at = ''
        field.checked_by = ''
    return field

def mark_checked(field, note, reviewer):
    field.approved = True
    field.review_note = note.strip()
    field.checked_at = dt.datetime.now(dt.timezone.utc).isoformat()
    field.checked_by = reviewer
    field.review_fingerprint = field_fingerprint(field)

def approved(field):
    return bool(field.approved and field.review_note.strip() and field.checked_by
                and field.checked_at and field.review_fingerprint == field_fingerprint(field))

def validate_rfq(rfq, publish=True):
    errors=[]
    if not rfq.title.strip(): errors.append('Give your request a title.')
    if len(rfq.items)>MAX_ITEMS: errors.append(f'At most {MAX_ITEMS} items are supported.')
    for name, fields in [('Item',rfq.items),('Question',rfq.questions)]:
        ids=[x.id.strip().casefold() for x in fields]
        if any(not x for x in ids) or len(set(ids))!=len(ids):
            errors.append(f'{name} IDs must be unique, including case and whitespace.')
    if publish and not rfq.items: errors.append('Add at least one item.')
    for item in rfq.items:
        if not item.description.strip() or not item.uom.strip(): errors.append(f'{item.id}: description and unit are required.')
        if publish and item.quantity is None: errors.append(f'{item.id}: supply a quantity before finishing.')
    for q in rfq.questions:
        if not q.label.strip(): errors.append(f'{q.id}: write a supplier question.')
        if q.mandatory and q.rule=='information': errors.append(f'{q.id}: required questions need a pass rule.')
        if q.rule in {'min','max'} and (q.threshold is None or not q.unit.strip()): errors.append(f'{q.id}: define a threshold and unit.')
    for term in rfq.terms:
        if not term.name.strip() or not term.value.strip(): errors.append('Each term needs a name and details.')
    if publish and rfq.open_questions: errors.append('Resolve the outstanding requirement questions before finishing.')
    return errors

def normalize(bid, item, fx, today=None, cost_mode='goods'):
    core=c(); today=today or dt.date.today()
    def stop(reason, status='NOT COMPARABLE'):
        return {'status':status,'price_inr':None,'normalized_decimal':None,'steps':'','reason':reason}
    if bid.excluded: return stop('Excluded alternative quote.','EXCLUDED')
    price=core.number(bid.price)
    if price is None or price<=0: return stop('No usable price.','MISSING')
    curr=bid.currency.strip().upper(); rate=core.number(fx.get(curr)); basis=core.number(bid.price_basis)
    if rate is None or rate<=0: return stop(f'No usable planning FX rate for {curr or "unknown currency"}.')
    if curr=='INR' and rate!=1: return stop('INR must have a rate of exactly 1.')
    if basis is None or basis<=0: return stop('Price basis is missing.')
    source,target=core.unit(bid.quoted_uom),core.unit(item.uom)
    if source and source==target: factor=Decimal(1)
    elif (source,target)==('ton','kg'): factor=Decimal(1000)
    elif (source,target)==('kg','ton'): factor=Decimal('.001')
    elif bid.target_units_per_quoted_unit and bid.conversion_evidence.strip(): factor=decimal(bid.target_units_per_quoted_unit)
    else: return stop(f'An evidenced conversion from {source or "unknown unit"} to {target} is required.')
    applied_discount=decimal(bid.discount_pct)
    if bid.discount_pct and not bid.discount_evidence.strip(): return stop('Discount needs source evidence.')
    if bid.discount_conditional:
        if bid.unconditional_base_price is None: return stop('Conditional discount needs an explicit unconditional base price.')
        price=decimal(bid.unconditional_base_price)
        applied_discount=Decimal(0)
        # Only an evidenced per-line quantity condition is executable. Whole-order,
        # supplier-spend and early-payment conditions retain the verified base price.
        if bid.conditional_min_quantity is not None:
            if not bid.constraints_checked or not bid.discount_evidence.strip(): return stop('Verify the line-quantity discount condition.')
            if decimal(item.quantity or 0)>=decimal(bid.conditional_min_quantity): applied_discount=decimal(bid.discount_pct)
    if bid.valid_until:
        try:
            if dt.date.fromisoformat(bid.valid_until)<today: return stop('Quote validity has expired.')
        except ValueError: return stop('Use YYYY-MM-DD for quote validity.')
    quantity=core.number(item.quantity)
    if quantity is None or quantity<=0: return stop('Requested quantity is missing.')
    if bid.capacity is not None and quantity>decimal(bid.capacity): return stop('Offered capacity is below the requested quantity.')
    if bid.min_order is not None and quantity<decimal(bid.min_order): return stop('Requested quantity is below the minimum order.')
    if bid.order_multiple is not None and quantity % decimal(bid.order_multiple): return stop('Requested quantity is not an allowed order multiple.')
    if item.max_lead_days is not None and (bid.lead_days is None or bid.lead_days>item.max_lead_days): return stop('Delivery lead time does not meet the requested limit.')
    if (bid.capacity is not None or bid.min_order is not None or bid.order_multiple is not None or bid.lead_days is not None or bid.valid_until) and not bid.constraints_checked:
        return stop('Verify the quoted validity, capacity and order constraints.','REVIEW REQUIRED')
    result=price*rate/basis/factor*(1-applied_discount/100)
    steps=f'{curr} {price} ÷ {basis} ÷ {factor} × FX {rate} × discount {1-applied_discount/100}'
    reason=''
    if not bid.file or not bid.locator.strip() or not bid.excerpt.strip(): reason='Missing source reference.'
    elif bid.uncertainty.strip() and not bid.issue_resolution.strip(): reason='Resolve the recorded uncertainty before approving.'
    elif not bid.specification_checked: reason='Confirm item matching and specifications.'
    elif not approved(bid): reason='This version of the price has not been checked.'
    elif cost_mode=='landed' and not bid.costs_checked: reason='Confirm tax, duty and line freight, including explicit zero amounts.'
    status='REVIEW REQUIRED' if reason else ('NORMALIZED' if curr!='INR' or basis!=1 or factor!=1 or bid.discount_pct else 'CONFIRMED')
    return {'status':status,'price_inr':float(result),'normalized_decimal':str(result),'steps':steps,'reason':reason}

def qualify(rfq, extraction, today=None):
    today=today or dt.date.today(); failures=[]; pending=[]
    for q in rfq.questions:
        if not q.mandatory: continue
        matches=[a for a in extraction.answers if a.question_id==q.id]
        if len(matches)!=1: pending.append(f'{q.label}: missing or conflicting answers'); continue
        a=matches[0]
        if not approved(a) or not a.file or not a.locator.strip() or not a.excerpt.strip():
            pending.append(f'{q.label}: verify this version of the answer'); continue
        if q.rule=='valid_certificate':
            if a.certificate_state in {'NO','EXPIRED'}: failures.append(f'{q.label}: {a.certificate_state.lower()}')
            elif a.certificate_state!='VALID': pending.append(f'{q.label}: a valid certificate is required; N/A does not waive this rule')
            else:
                try:
                    if dt.date.fromisoformat(a.expiry)<today: failures.append(f'{q.label}: expired {a.expiry}')
                except ValueError: pending.append(f'{q.label}: a valid expiry date is required')
        elif q.rule=='yes':
            value=a.value.strip().lower()
            if value in {'no','false'}: failures.append(f'{q.label}: no')
            elif value not in {'yes','true'}: pending.append(f'{q.label}: yes/no is unclear')
        elif q.rule in {'min','max'}:
            if a.numeric_value is None or c().unit(a.unit)!=c().unit(q.unit): pending.append(f'{q.label}: missing value or incompatible unit')
            elif q.threshold is None: pending.append(f'{q.label}: missing buyer threshold')
            elif (q.rule=='max' and a.numeric_value>q.threshold) or (q.rule=='min' and a.numeric_value<q.threshold): failures.append(f'{q.label}: threshold failed')
    state='DISQUALIFIED' if failures else 'PENDING' if pending else 'QUALIFIED'
    return {'state':state,'reasons':failures+pending or ['All defined required rules passed.' if any(q.mandatory for q in rfq.questions) else 'No required qualification rules were defined.']}

def build_dataset(rfq,responses,fx):
    rows=[]; questions=[]; commercials=[]
    for supplier,response in responses.items():
        ext=c().Extraction.model_validate(response['extraction']); qualification=qualify(rfq,ext)
        questions.append({'supplier':supplier,**qualification})
        policy=response.get('policy',{})
        mode=policy.get('cost_mode','goods')
        for item in rfq.items:
            matches=[b for b in ext.bids if b.line_id==item.id and not b.excluded]
            if len(matches)==1:
                bid=matches[0]; n=normalize(bid,item,fx,cost_mode=mode)
            else:
                bid=c().Bid(line_id=item.id)
                n={'status':'CONFLICT' if matches else 'MISSING','price_inr':None,'normalized_decimal':None,'steps':'','reason':'Choose one current quote and exclude the alternatives.' if matches else 'No current quote for this item.'}
            if mode=='landed' and not policy.get('fixed_charge_checked') and n['status'] in c().STATUS_OK:
                n={**n,'status':'REVIEW REQUIRED','reason':'Confirm supplier-wide fixed charges, including explicit zero.'}
            rows.append({'evidence_id':f'{supplier}|{item.id}','supplier':supplier,'description':item.description,
                'specification':item.specification,'quantity':item.quantity,'uom':item.uom,'location':item.location,
                **c().dump(bid),**n,'qualification':qualification['state'],
                'eligible':n['status'] in c().STATUS_OK and qualification['state']=='QUALIFIED',
                'cost_mode':mode,'supplier_fixed_inr':policy.get('fixed_charge_inr',0),
                'supplier_costs_checked':policy.get('fixed_charge_checked',False)})
        for idx,a in enumerate(ext.answers): questions.append({'evidence_id':f'{supplier}|Q|{idx}','supplier':supplier,**c().dump(a)})
        for idx,t in enumerate(ext.commercials): commercials.append({'evidence_id':f'{supplier}|C|{idx}','supplier':supplier,**c().dump(t),'terms_checked':response.get('terms_checked',False)})
    return rows,questions,commercials

def line_cost(row,item,cost_mode):
    goods=cents(decimal(row.get('normalized_decimal') or row['price_inr'])*decimal(item.quantity))
    extra=Decimal(0)
    if cost_mode=='landed':
        extra=cents(goods*(decimal(row.get('tax_pct',0))+decimal(row.get('duty_pct',0)))/100)+cents(row.get('freight_inr',0))
    return goods,cents(goods+extra)

def allocate(items,rows,suppliers,max_suppliers=None,cost_mode='goods'):
    """Exact DP over supplier masks; a line is never fractionally split."""
    if len(suppliers)>8 or len(set(suppliers))!=len(suppliers): raise ValueError('Use at most eight unique suppliers.')
    if len(items)>MAX_ITEMS: raise ValueError('Item limit exceeded.')
    cap=max_suppliers or len(suppliers)
    if cap<1 or cap>8: raise ValueError('Supplier cap must be between 1 and 8.')
    bits={s:1<<i for i,s in enumerate(sorted(suppliers))}
    fixed={s:Decimal(0) for s in suppliers}
    if cost_mode=='landed':
        for s in suppliers:
            rs=[r for r in rows if r['supplier']==s]
            charges={cents(r.get('supplier_fixed_inr',0)) for r in rs}
            if len(charges)>1: raise ValueError('Inconsistent fixed supplier charges.')
            fixed[s]=next(iter(charges),Decimal(0))
    states={0:(Decimal(0),[])}; missing=[]
    for item in items:
        options=[r for r in rows if r['eligible'] and r['line_id']==item.id and r['supplier'] in bits
                 and (cost_mode!='landed' or (r.get('costs_checked') and r.get('supplier_costs_checked')))]
        if not options:
            missing.append({'line_id':item.id,'description':item.description,'quantity':item.quantity,'uom':item.uom,'supplier':'UNASSIGNED','price_inr':None,'spend_inr':None,'goods_spend_inr':None,'fixed_charge_inr':0,'evidence_id':'','why':'No checked price meeting the requested rules.'})
            continue
        new={}
        for mask,(total,allocations) in states.items():
            for row in sorted(options,key=lambda r:r['supplier']):
                next_mask=mask|bits[row['supplier']]
                if next_mask.bit_count()>cap: continue
                goods,spend=line_cost(row,item,cost_mode)
                charge=fixed[row['supplier']] if not mask&bits[row['supplier']] else Decimal(0)
                entry={'line_id':item.id,'description':item.description,'quantity':item.quantity,'uom':item.uom,
                    'supplier':row['supplier'],'price_inr':row['price_inr'],'goods_spend_inr':float(goods),
                    'fixed_charge_inr':float(charge),'spend_inr':float(spend+charge),'evidence_id':row['evidence_id'],
                    'why':'Checked source, accepted specifications, required rules passed; lowest total under the selected constraints.'}
                candidate=(total+spend+charge,allocations+[entry])
                old=new.get(next_mask)
                if old is None or (candidate[0],tuple(x['supplier'] for x in candidate[1]))<(old[0],tuple(x['supplier'] for x in old[1])): new[next_mask]=candidate
        states=new
        if not states:
            return {'spend_inr':None,'covered':0,'total_lines':len(items),'complete':False,'supplier_count':None,'allocation':[], 'cost_mode':cost_mode,'reason':'No complete allocation meets the supplier cap.'}
    if not states: states={0:(Decimal(0),[])}
    mask,(total,allocation)=min(states.items(),key=lambda x:(x[1][0],x[0].bit_count(),tuple(a['supplier'] for a in x[1][1])))
    lookup={a['line_id']:a for a in allocation+missing}
    allocation=[lookup[i.id] for i in items]
    return {'spend_inr':float(total),'covered':len(items)-len(missing),'total_lines':len(items),
        'complete':bool(items) and not missing,'supplier_count':mask.bit_count(),'allocation':allocation,'cost_mode':cost_mode}

def scenario_engine(items,rows,suppliers,max_suppliers=None,cost_mode=None):
    mode=cost_mode or next((r.get('cost_mode','goods') for r in rows),'goods')
    cap=max_suppliers or len(suppliers)
    if not suppliers: return [{'strategy':'Lowest eligible cost',**allocate(items,rows,['__none__'],cost_mode=mode)}]
    unrestricted=allocate(items,rows,suppliers,cost_mode=mode)
    scenarios=[{'strategy':'Lowest eligible cost',**unrestricted}]
    for k in sorted({1,min(2,len(suppliers)),min(cap,len(suppliers))}):
        result=allocate(items,rows,suppliers,k,mode)
        if not result['complete']:
            result={**result,'spend_inr':None,'allocation':[]}
        scenarios.append({'strategy':f'At most {k} supplier(s)',**result})
    single=next((s for s in scenarios if s['strategy']=='At most 1 supplier(s)' and s['complete']),None)
    for s in scenarios: s['savings_vs_single_inr']=float(cents(decimal(single['spend_inr'])-decimal(s['spend_inr']))) if single and s['complete'] else None
    return scenarios

def safe_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        members=z.infolist()
        if len(members)>5000 or sum(m.file_size for m in members)>MAX_EXPANDED: raise ValueError('Compressed document expands beyond the safe limit.')
        for m in members:
            if m.flag_bits&1: raise ValueError('Encrypted documents are unsupported.')
            if m.file_size>20*1024*1024: raise ValueError('An embedded document part exceeds the limit.')
            if m.file_size>1024*1024 and m.file_size/max(1,m.compress_size)>200: raise ValueError('Unusually compressed document rejected.')
    return True

def source_text(filename,data):
    if not data or len(data)>c().MAX_FILE: raise ValueError('Document is empty or exceeds 18 MB.')
    ext=Path(filename).suffix.lower()
    if ext in {'.xlsx','.docx'}: safe_archive(data)
    if ext=='.pdf':
        import fitz
        with fitz.open(stream=data,filetype='pdf') as doc:
            if doc.is_encrypted or not 0<len(doc)<=MAX_PAGES: raise ValueError('Use an unencrypted PDF with 1–300 pages.')
        return None
    if ext in {'.png','.jpg','.jpeg'}:
        from PIL import Image
        with Image.open(io.BytesIO(data)) as image:
            if image.width*image.height>24_000_000: raise ValueError('Image exceeds the 24-megapixel limit.')
            image.verify()
        return None
    if ext=='.xlsx':
        import openpyxl
        wb=openpyxl.load_workbook(io.BytesIO(data),data_only=True,read_only=True)
        out=[]; count=0
        try:
            for sheet in wb:
                if sheet.max_row>10000 or sheet.max_column>100: raise ValueError('Spreadsheet dimensions exceed 10,000 rows / 100 columns.')
                for idx,values in enumerate(sheet.iter_rows(values_only=True),1):
                    cells=' | '.join(f'{openpyxl.utils.get_column_letter(i)}={v}' for i,v in enumerate(values,1) if v is not None)
                    if cells: out.append(f'Sheet {sheet.title}, Row {idx}: {cells}'); count+=len(out[-1])
                    if count>MAX_TEXT: raise ValueError('Spreadsheet text is too large; split the document.')
        finally: wb.close()
        # Do not silently lose formulas with no cached values.
        formulas=openpyxl.load_workbook(io.BytesIO(data),data_only=False,read_only=True)
        try:
            if any(cell.data_type=='f' for s in formulas for row in s.iter_rows() for cell in row):
                out.insert(0,'READER WARNING: formulas use saved cached values only; recalculate and save in Excel before uploading if results are blank or stale.')
        finally: formulas.close()
        return '\n'.join(out)
    if ext=='.eml':
        from email import policy
        from email.parser import BytesParser
        from html.parser import HTMLParser
        message=BytesParser(policy=policy.default).parsebytes(data)
        if any(p.get_content_disposition()=='attachment' for p in message.walk()): raise ValueError('Upload email attachments separately alongside the email.')
        body=message.get_body(preferencelist=('plain','html'))
        text=str(body.get_content()) if body else str(message.get_content())
        if body and body.get_content_type()=='text/html':
            class TextReader(HTMLParser):
                def __init__(self): super().__init__(); self.parts=[]; self.hidden=0
                def handle_starttag(self,tag,attrs):
                    if tag in {'script','style'}: self.hidden+=1
                def handle_endtag(self,tag):
                    if tag in {'script','style'}: self.hidden=max(0,self.hidden-1)
                def handle_data(self,value):
                    if not self.hidden: self.parts.append(value)
            reader=TextReader(); reader.feed(text); text=' '.join(reader.parts)
        return '\n'.join(f'Line {i}: {line}' for i,line in enumerate(text.splitlines(),1))
    result=c()._source_text_impl(filename,data)
    if result is not None and len(result)>MAX_TEXT: raise ValueError('Document text exceeds 250,000 characters.')
    return result

def prepare_sources(files):
    return c()._prepare_sources_impl(files)

def verify_evidence(field,files):
    file=next((f for f in files if f['name']==field.file),None)
    if not file or not field.locator.strip() or not field.excerpt.strip(): return False,'Add a valid file, locator and excerpt.'
    raw=base64.b64decode(file['data'],validate=True) if isinstance(file['data'],str) else file['data']
    ext=Path(file['name']).suffix.lower()
    if ext=='.pdf':
        import fitz
        match=re.fullmatch(r'Page\s+(\d+)',field.locator.strip(),re.I)
        with fitz.open(stream=raw,filetype='pdf') as doc:
            if not match or not 1<=int(match[1])<=len(doc): return False,'Reference an existing PDF page as Page 1.'
        return True,'Manual source verification required.'
    if ext in {'.png','.jpg','.jpeg'}: return True,'Manual source verification required.'
    text=source_text(file['name'],raw)
    locator=field.locator.strip()
    if Path(file['name']).suffix.lower()=='.xlsx':
        match=re.fullmatch(r'Sheet\s+(.+?),\s*Row\s+(\d+)',locator,re.I)
        lines=[line for line in text.splitlines() if match and re.match(r'Sheet\s+'+re.escape(match[1])+r',\s*Row\s+'+match[2]+r':',line,re.I)]
        lines=[re.sub(r'(^|\s*\|\s*)[A-Z]+=',lambda m:m[1],line.split(':',1)[-1]) for line in lines]
    else:
        lines=[line.split(':',1)[-1] for line in text.splitlines() if line.lower().startswith(locator.lower()+':')]
    if not lines: return False,'The stated locator was not found in the document.'
    if not any(c().squash(field.excerpt) in c().squash(line) for line in lines): return False,'The excerpt does not match the stated locator.'
    return True,''

def approve_review(bids,answers,known,names,confirm_all,note,files=None,reviewer='Local buyer'):
    counts={i:sum(b.line_id==i and not b.excluded for b in bids) for i in known}
    acounts={a.question_id:sum(x.question_id==a.question_id for x in answers) for a in answers}
    problems=[]
    for field in list(bids)+list(answers):
        old_requested=field.approved
        changed=bool(field.review_fingerprint and field.review_fingerprint!=field_fingerprint(field))
        invalidate_changed(field)
        requested=confirm_all or old_requested
        if not requested: continue
        if isinstance(field,c().Bid):
            valid=(field.line_id in known and counts.get(field.line_id)==1 and not field.excluded and field.price
                   and field.price_basis and field.currency.strip() and field.quoted_uom.strip() and field.specification_checked
                   and (not field.uncertainty.strip() or field.issue_resolution.strip())
                   and (not field.discount_conditional or field.unconditional_base_price is not None)
                   and (not field.discount_pct or field.discount_evidence.strip()))
        else:
            valid=acounts.get(field.question_id)==1 and (field.value.strip() or field.numeric_value is not None or field.certificate_state not in {'MISSING','CLAIMED'})
        valid=valid and field.file in names and field.locator.strip() and field.excerpt.strip()
        if files:
            supported,why=verify_evidence(field,files)
            valid=valid and supported
            if not supported: problems.append(why)
        # A changed previously-approved value requires an explicit new confirmation.
        if changed and not confirm_all: valid=False
        if valid and requested and note.strip(): mark_checked(field,note,reviewer)
        else: field.approved=False; field.review_fingerprint=''
    return list(dict.fromkeys(problems))

class Strict(BaseModel):
    model_config=ConfigDict(extra='forbid',allow_inf_nan=False)

class StoredFile(Strict):
    name:str=Field(min_length=1,max_length=200)
    data:str
    hash:str

class Policy(Strict):
    cost_mode:str='goods'
    fixed_charge_inr:float=Field(default=0,ge=0)
    fixed_charge_checked:bool=False

class Response(Strict):
    supplier:str
    rfq_hash:str
    review_id:str=''
    extraction:dict
    bundle_hash:str=''
    files:list[StoredFile]=Field(max_length=12)
    terms_checked:bool=False
    policy:Policy=Field(default_factory=Policy)
    revision:int=Field(default=1,ge=1)

class Workspace(Strict):
    version:int=2
    workspace_id:str=''
    rfq:dict|None=None
    published:bool=False
    suppliers:list[str]=Field(default_factory=list,max_length=8)
    responses:dict[str,Response]=Field(default_factory=dict)
    fx:dict[str,float]=Field(default_factory=lambda:{'INR':1.0})
    fx_date:str=Field(default_factory=lambda:dt.date.today().isoformat())
    fx_basis:str='INR base currency'
    fx_reference:dict=Field(default_factory=dict)
    events:list[dict]=Field(default_factory=list,max_length=5000)
    conversation:list[dict]=Field(default_factory=list,max_length=1000)
    demo_invited_suppliers:list[str]=Field(default_factory=list)
    quote_versions:dict[str,list[dict]]=Field(default_factory=dict)
    rfq_versions:list[dict]=Field(default_factory=list)
    final_selection:dict|None=None
    pending_review:dict|None=None
    cost_mode:str='goods'
    clarifications:list[dict]=Field(default_factory=list)
    request_sources:list[StoredFile]=Field(default_factory=list,max_length=12)

def empty_workspace():
    import uuid
    return Workspace(workspace_id=uuid.uuid4().hex).model_dump()

def validate_response(response,rfq,supplier):
    if response.supplier!=supplier: raise ValueError('Response supplier does not match its workspace key.')
    ext=c().Extraction.model_validate(response.extraction)
    known={i.id for i in rfq.items}; qids={q.id for q in rfq.questions}
    if any(b.line_id is not None and b.line_id not in known for b in ext.bids): raise ValueError('Unknown request item in a response.')
    if any(a.question_id not in qids for a in ext.answers): raise ValueError('Unknown question in a response.')
    names=[f.name for f in response.files]
    if len(set(names))!=len(names): raise ValueError('Duplicate source filenames.')
    size=0
    for file in response.files:
        if Path(file.name).name!=file.name or '/' in file.name or '\\' in file.name: raise ValueError('Source filenames must not include paths.')
        raw=base64.b64decode(file.data,validate=True); size+=len(raw)
        if len(raw)>c().MAX_FILE or hashlib.sha256(raw).hexdigest()!=file.hash: raise ValueError('Invalid file size or source hash.')
        source_text(file.name,raw)
    if size>c().MAX_BUNDLE: raise ValueError('Source bundle exceeds 36 MB.')
    for field in ext.bids+ext.answers+ext.commercials:
        if field.file and field.file not in names: raise ValueError('A response references an unknown source file.')
    if response.policy.cost_mode not in {'goods','landed'}: raise ValueError('Invalid cost mode.')
    return ext

def load_workspace(data,trusted_local=False):
    if len(data)>100*1024*1024: raise ValueError('Backup exceeds 100 MB.')
    obj=json.loads(data,parse_constant=lambda x:(_ for _ in ()).throw(ValueError('Nonfinite number in backup.')))
    if not isinstance(obj,dict) or obj.get('version') not in {1,2}: raise ValueError('Unsupported backup version.')
    obj={**obj,'version':2}
    workspace=Workspace.model_validate(obj)
    if workspace.cost_mode not in {'goods','landed'}: raise ValueError('Invalid cost mode.')
    if any(not s.strip() for s in workspace.suppliers) or len({s.strip().casefold() for s in workspace.suppliers})!=len(workspace.suppliers): raise ValueError('Supplier names must be nonempty and unique.')
    if set(workspace.responses)-set(workspace.suppliers): raise ValueError('Response exists for a supplier outside the list.')
    if workspace.fx.get('INR')!=1 or any(not re.fullmatch('[A-Z]{3}',key) or value<=0 for key,value in workspace.fx.items()): raise ValueError('FX rates must be positive, with INR fixed at 1.')
    dt.date.fromisoformat(workspace.fx_date)
    for currency,ref in workspace.fx_reference.items():
        if currency not in workspace.fx: raise ValueError('Unknown FX reference currency.')
        date=dt.date.fromisoformat(ref['date'])
        if date>dt.date.today(): raise ValueError('FX reference date is in the future.')
        if decimal(ref['rate'])!=decimal(workspace.fx[currency]): raise ValueError('FX reference and rate differ.')
    rfq=c().RFQ.model_validate(workspace.rfq) if workspace.rfq else None
    if workspace.published and (not rfq or validate_rfq(rfq)): raise ValueError('Published request is invalid or has unresolved questions.')
    if workspace.responses and (not workspace.published or rfq is None): raise ValueError('Quotes require a finished request.')
    for supplier,response in workspace.responses.items():
        ext=validate_response(response,rfq,supplier)
        if not trusted_local:
            for field in ext.bids+ext.answers: field.approved=False; field.review_fingerprint=''; field.checked_at=''; field.checked_by=''
            response.terms_checked=False; response.policy.fixed_charge_checked=False
        response.extraction=c().dump(ext)
    if workspace.pending_review:
        pending=Response.model_validate(workspace.pending_review)
        if not rfq or pending.supplier not in workspace.suppliers: raise ValueError('Pending review is outside the current request.')
        pending_ext=validate_response(pending,rfq,pending.supplier)
        if not trusted_local:
            for field in pending_ext.bids+pending_ext.answers: field.approved=False; field.review_fingerprint=''; field.checked_at=''; field.checked_by=''
            pending.extraction=c().dump(pending_ext); pending.terms_checked=False; pending.policy.fixed_charge_checked=False
        workspace.pending_review=pending.model_dump()
    for supplier,versions in workspace.quote_versions.items():
        if supplier not in workspace.suppliers or len(versions)>10: raise ValueError('Invalid quote revision history.')
        for old in versions:
            revision=Response.model_validate(old)
            validate_response(revision,rfq,supplier)
    for old in workspace.rfq_versions:
        c().RFQ.model_validate(old['rfq'])
    request_size=0
    for file in workspace.request_sources:
        raw=base64.b64decode(file.data,validate=True); request_size+=len(raw)
        if hashlib.sha256(raw).hexdigest()!=file.hash: raise ValueError('Request document hash mismatch.')
        source_text(file.name,raw)
    if request_size>c().MAX_BUNDLE: raise ValueError('Requirement documents exceed the bundle limit.')
    # History is archival, never an input to calculations. Clear untrusted decisions.
    if not trusted_local:
        workspace.final_selection=None; workspace.conversation=[]
    result=workspace.model_dump()
    if not result['workspace_id']:
        import uuid
        result['workspace_id']=uuid.uuid4().hex
    return result

def workspace_bytes(workspace):
    payload=json.dumps(workspace,ensure_ascii=False,indent=2,allow_nan=False).encode('utf-8')
    if len(payload)>100*1024*1024: raise ValueError('Workspace exceeds 100 MB. Start a separate request or archive older large quote bundles.')
    return payload

def record(workspace,action,details):
    workspace.setdefault('events',[]).append({'time_utc':dt.datetime.now(dt.timezone.utc).isoformat(),'action':action,'details':details})
    workspace['events']=workspace['events'][-5000:]

def rate_is_fresh(workspace,currency,today=None):
    today=today or dt.date.today()
    if currency=='INR': return workspace.get('fx',{}).get('INR')==1
    ref=workspace.get('fx_reference',{}).get(currency)
    if not ref: return False
    try: age=(today-dt.date.fromisoformat(ref['date'])).days
    except (ValueError,KeyError): return False
    return 0<=age<=10 and workspace.get('fx',{}).get(currency,0)>0

def effective_fx(workspace):
    return {key:value for key,value in workspace['fx'].items() if rate_is_fresh(workspace,key)}

def fetch_reference_rate(currency):
    return c()._fetch_reference_rate_impl(currency)
